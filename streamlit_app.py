"""
Streamlit web version of the Order Quotation app.

This file is the one to deploy on Streamlit Cloud - point your app's
"Main file path" at streamlit_app.py (not app.py). All the pricing /
discount / matching logic lives in app.py and is imported here unchanged,
so any CONFIG tweaks you make in app.py (discount %, extra rates, etc.)
apply here too.
"""

import io
import os
import tempfile

import streamlit as st
import openpyxl

from app import (
    load_rate_list,
    compute_quotation_row,
    generate_quotation_excel,
    extract_items_from_pdf,
    update_rate_list_price,
    find_column,
)

st.set_page_config(page_title="Pipe & Fitting Order Quotation", layout="wide")
st.title("Pipe & Fitting Order Quotation")


# ---------------------------------------------------------------------
# Rate list upload (shared across all tabs)
# ---------------------------------------------------------------------

st.subheader("1. Upload your master rate list")
rate_file = st.file_uploader(
    "RATE_LIST_PIP___FITTING.xlsx (or any workbook with the same PVC (2) / PVC / APVC / CPVC sheets)",
    type=["xlsx"],
)

if rate_file is not None:
    st.session_state["rate_bytes"] = rate_file.getvalue()
    st.session_state["rate_name"] = rate_file.name


@st.cache_data(show_spinner="Loading rate list...")
def _load_tables(rate_bytes):
    return load_rate_list(io.BytesIO(rate_bytes))


if "rate_bytes" not in st.session_state:
    st.info("Upload your rate list above to get started.")
    st.stop()

rate_tables = _load_tables(st.session_state["rate_bytes"])
st.success(f"Rate list loaded: {st.session_state['rate_name']}  "
           f"({sum(len(t) for t in rate_tables.values())} items across "
           f"{len(rate_tables)} categories)")


# ---------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------

def rows_to_xlsx_bytes(rows):
    """generate_quotation_excel() writes to a path, so we hand it a temp
    file and read the bytes back - keeps app.py identical for the
    command-line version."""
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    try:
        generate_quotation_excel(rows, tmp.name)
        with open(tmp.name, "rb") as f:
            return f.read()
    finally:
        os.unlink(tmp.name)


def show_rows_table(rows):
    import pandas as pd
    df = pd.DataFrame(rows)[
        ["name", "family", "order_qty", "billing_qty", "rate",
         "discount_pct", "rate_after_discount", "amount", "note"]
    ]
    df.columns = ["Product Name", "Family", "Order Qty", "Billing Qty",
                  "Rate", "Discount %", "Rate After Discount", "Amount", "Note"]
    flagged = df["Note"].str.contains("LOW CONFIDENCE|NOT FOUND", na=False)
    if flagged.any():
        st.warning(f"{flagged.sum()} item(s) need a manual check - see the Note column.")
    st.dataframe(df, width='stretch')
    st.metric("Grand Total", f"{df['Amount'].sum():,.2f}")


# ---------------------------------------------------------------------
# Tabs
# ---------------------------------------------------------------------

tab1, tab2, tab3 = st.tabs([
    "Generate Order Quotation",
    "Extract items from a Sales Bill PDF",
    "Update a rate list price",
])

# --- Tab 1: Generate Order Quotation ---------------------------------
with tab1:
    st.write("Enter one item per line, as:  `Item Name, Qty`")
    default_text = (
        "PVC PIPE 1I X 6KG ISI AST, 100\n"
        "PVC ELBOW 1I, 10\n"
        "PVC TEE 1I, 5\n"
    )
    items_text = st.text_area("Items", value=default_text, height=200, key="items_text")

    uploaded_csv = st.file_uploader("...or upload a CSV instead (Item Name, Qty columns)",
                                     type=["csv"], key="csv_upload")

    if st.button("Generate Quotation", type="primary"):
        items = []
        if uploaded_csv is not None:
            import csv as csv_module
            text = uploaded_csv.getvalue().decode("utf-8-sig")
            reader = csv_module.reader(io.StringIO(text))
            for row in reader:
                if not row or not row[0].strip():
                    continue
                if row[0].strip().upper() in ("ITEM", "PRODUCT NAME", "NAME"):
                    continue
                try:
                    items.append((row[0].strip(), float(row[1].strip())))
                except (IndexError, ValueError):
                    st.warning(f"Skipped row (couldn't read quantity): {row}")
        else:
            for line in items_text.splitlines():
                line = line.strip()
                if not line:
                    continue
                if "," not in line:
                    st.warning(f"Skipped line (no comma found): {line}")
                    continue
                name, qty_raw = line.rsplit(",", 1)
                try:
                    items.append((name.strip(), float(qty_raw.strip())))
                except ValueError:
                    st.warning(f"Skipped line (couldn't read quantity): {line}")

        if not items:
            st.error("No valid items found.")
        else:
            rows = [compute_quotation_row(name, qty, rate_tables) for name, qty in items]
            st.session_state["last_quotation_rows"] = rows
            show_rows_table(rows)
            xlsx_bytes = rows_to_xlsx_bytes(rows)
            st.download_button(
                "Download Order_Quotation.xlsx",
                data=xlsx_bytes,
                file_name="Order_Quotation.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# --- Tab 2: Extract from PDF ------------------------------------------
with tab2:
    pdf_file = st.file_uploader("Sales Bill PDF", type=["pdf"], key="pdf_upload")
    families = st.multiselect("Families to extract", ["PVC", "APVC", "CPVC"],
                               default=["PVC", "APVC", "CPVC"])

    if pdf_file is not None and st.button("Extract items"):
        items = extract_items_from_pdf(pdf_file, families=tuple(families))
        if not items:
            st.error("No matching items found - check the PDF layout / families chosen.")
        else:
            st.session_state["extracted_items"] = items
            st.success(f"Found {len(items)} item(s).")
            import pandas as pd
            st.dataframe(pd.DataFrame(items, columns=["Item Name", "Qty"]), width='stretch')

    if "extracted_items" in st.session_state and st.button("Fill rates & build quotation"):
        items = st.session_state["extracted_items"]
        rows = [compute_quotation_row(name, qty, rate_tables) for name, qty in items]
        show_rows_table(rows)
        xlsx_bytes = rows_to_xlsx_bytes(rows)
        st.download_button(
            "Download Extracted_Items.xlsx",
            data=xlsx_bytes,
            file_name="Extracted_Items.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

# --- Tab 3: Update a rate list price -----------------------------------
with tab3:
    wb_preview = openpyxl.load_workbook(io.BytesIO(st.session_state["rate_bytes"]), read_only=True)
    sheet_name = st.selectbox("Sheet", wb_preview.sheetnames)
    search_text = st.text_input("Text to search for in PRODUCT NAME (matches any row containing it)")
    column_name = st.selectbox("Column to update", ["RATE", "PER/RATE"])
    new_rate = st.number_input("New rate", min_value=0.0, step=0.5)

    if st.button("Preview matching rows"):
        wb_check = openpyxl.load_workbook(io.BytesIO(st.session_state["rate_bytes"]), data_only=True)
        ws = wb_check[sheet_name]
        rows_iter = list(ws.iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows_iter)
                            if any(v and str(v).strip().upper() == "PRODUCT NAME" for v in r)), None)
        if header_idx is None:
            st.error("Couldn't find the header row in this sheet.")
        else:
            header = rows_iter[header_idx]
            name_col = find_column(header, "PRODUCT NAME")
            matches = [r[name_col] for r in rows_iter[header_idx + 1:]
                       if r[name_col] and search_text.strip().upper() in str(r[name_col]).upper()]
            if matches:
                st.write(f"{len(matches)} row(s) will be updated to {column_name} = {new_rate}:")
                st.write(matches)
            else:
                st.warning("No matching rows found.")

    if st.button("Apply update & download new rate list", type="primary"):
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        tmp.write(st.session_state["rate_bytes"])
        tmp.close()
        try:
            changed = update_rate_list_price(tmp.name, sheet_name, search_text, new_rate,
                                              column_name=column_name)
            if changed:
                with open(tmp.name, "rb") as f:
                    updated_bytes = f.read()
                st.success(f"Updated {len(changed)} row(s).")
                st.download_button(
                    "Download updated rate list",
                    data=updated_bytes,
                    file_name=st.session_state["rate_name"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            else:
                st.warning("No matching rows found - nothing was changed.")
        finally:
            os.unlink(tmp.name)
