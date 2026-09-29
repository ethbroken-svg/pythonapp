#!/usr/bin/env python3
"""
Salman's Pipe & Fitting Order Quotation App
=============================================
Automates the order-quotation workflow used for the shop's PVC / APVC / CPVC
pipe & fitting business:

  1. Generate a priced, discounted Order Quotation (Excel) from a simple
     item list (name + quantity) - matches each item to the master rate
     list, applies the correct discount/qty rules, and writes a formatted
     .xlsx file.

  2. Extract PVC / APVC / CPVC items + quantities from a scanned Sales Bill
     (PDF), ignore the bill's own price, and fill in the correct rate from
     the master rate list.

  3. Update a price in the master rate list (single item).

Run it with:
    python3 app.py

Everything you might need to tweak (file paths, discount %, extra manual
rates, special no-discount keywords) is in the CONFIG section below - you
do not need to touch anything past "===== END CONFIG =====" for normal use.
"""

import os
import re
import sys
import csv
import glob
import difflib

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# =====================================================================
# Small text helper used by CONFIG below, so it's defined first
# =====================================================================

def normalize(text):
    """Uppercase, collapse whitespace, and drop punctuation noise so that
    e.g.  PVC 1II PIPE 6KG ISI AST   and   pvc 1ii  pipe 6kg isi ast
    compare equal, and so quote/inch marks don't break matching."""
    if text is None:
        return ""
    t = str(text).upper()
    t = t.replace('"', "").replace("''", "").replace("'", "")
    t = re.sub(r"[^A-Z0-9./ -]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    # join split weight markings like "10 KG" -> "10KG" so they tokenize
    # the same way the rate list does
    t = re.sub(r"\b(\d+)\s*KG\b", r"\1KG", t)
    return t


# =====================================================================
# CONFIG - edit this section as your business rules change
# =====================================================================

# Default path to the master rate list. If it's not found, the app will
# ask you to type the path (or look for any .xlsx in the current folder).
RATE_LIST_PATH = "RATE_LIST_PIP___FITTING.xlsx"

# Which sheet + which column ("RATE" or "PER/RATE") to read for each
# product family. Worked out by checking the actual workbook:
#   - PVC pipes: the "PVC" sheet's RATE column is a PER-FOOT price, so for
#     pipes we instead read the "PVC (2)" sheet's PER/RATE column, which
#     holds the full-pipe price.
#   - PVC fittings: the "PVC" sheet's RATE column already holds the full
#     item price.
#   - APVC: the "APVC" sheet's RATE column holds the full price for both
#     pipes and fittings.
#   - CPVC: the "CPVC" sheet's RATE column is the correct one to bill from
#     (confirmed against Salman's own examples: CPVC OIII PIPE AST = 445,
#     COUPLER = 25, TEE = 46, ELBOW = 26 - all come from RATE, not the
#     lower PER/RATE column).
SHEET_RULES = {
    "PVC_PIPE":    {"sheet": "PVC (2)", "column": "PER/RATE"},
    "PVC_FITTING": {"sheet": "PVC",     "column": "RATE"},
    "APVC":        {"sheet": "APVC",    "column": "RATE"},
    "CPVC":        {"sheet": "CPVC",    "column": "RATE"},
}

# Discount % by category (as a fraction, e.g. 0.37461 = 37.461%)
DISCOUNT_PVC_APVC_PIPE = 0.37461
DISCOUNT_PVC_APVC_FITTING = 0.3039
DISCOUNT_CPVC_PIPE = 0.445
DISCOUNT_CPVC_FITTING = 0.4218

# Pipe order-quantity divisors (order qty is assumed to be in FEET; this
# converts it to number of pipe pieces for billing)
PIPE_QTY_DIVISOR = {"PVC": 20, "APVC": 10, "CPVC": 10}

# Item-name keywords that NEVER get a discount, however they're matched.
# Matched as a substring against the normalized item name.
NO_DISCOUNT_KEYWORDS = [
    "SOLUTION", "LONG PLUG", "KHILLI", "KHILA", "KHILI",
    "CONCEAL VALVE", "TEFLON",
]

# Pipes that do NOT contain this brand word get no discount at all
# (fittings are unaffected by brand - they always get the category
# discount above). "10KG" / "NON ISI" pipes are also always no-discount
# (and see the special rate x20 rule below).
PIPE_DISCOUNT_REQUIRES_BRAND = "AST"

# PVC 10KG (non-ISI) pipes: the rate list stores a PER-FOOT price for
# these. Multiply that rate by 20 to get the equivalent full-pipe price
# (quantity is still divided by 20 as normal for PVC pipes). Net effect:
# billing exactly matches qty(feet) x per-foot-rate, matching how these
# actually show up on bills.
TEN_KG_RATE_MULTIPLIER = 20
TEN_KG_KEYWORDS = ["10KG", "10 KG"]

# Extra rates Salman has supplied by hand because they are not (or not
# correctly) in the master rate list.
# Key: normalize(item name as Salman writes it) -> (rate, counts_as_fitting, note)
# counts_as_fitting=True means the normal fitting discount for that family
# still applies; False means no discount at all.
EXTRA_RATES = {
    normalize("CPVC SOLUTION 20ML"): (40, False, "no discount (Solution) - rate given by Salman"),
    normalize("JASDI KHILA"): (120, False, "no discount (Khilli) - rate given by Salman"),
    normalize("PVC REDUCER 3 X 2II"): (59, True, "rate given by Salman"),
    normalize("PVC BEND 1II"): (65, True, "rate given by Salman"),
    normalize("PVC MULTI FLOOR GALLY TAP 4''"): (221, True, "same as 2II multi floor trap - rate given by Salman"),
}

# Known abbreviations / shorthand seen in bills, mapped to the word used
# in the rate list, so matching finds the right row. Applied as whole-word
# substitutions during normalization for matching purposes only.
SYNONYMS = {
    "MTA": "MALE",
    "CUPLIN": "COUPLER",
    "CUPLINE": "COUPLER",
    "CUPLING": "COUPLER",
    "ENDCUP": "END CAP",
    "GALLY TAP": "GALLY TRAP",
    "BRASS FEMALE": "BRASS FTA",
}

# Below this similarity score (0-1, token overlap) a match is flagged as
# LOW CONFIDENCE in the output "note" column, for you to double check.
# Always eyeball anything flagged - automatic matching can be fooled by
# oddly-worded bill entries (short codes, typos, reordered words).
CONFIDENCE_THRESHOLD = 0.65

# ===== END CONFIG =====


# =====================================================================
# Rate list loading
# =====================================================================

def find_column(header_row, name):
    """Return the 0-based index of the column whose header equals `name`
    (case-insensitive), or None if not found."""
    for i, h in enumerate(header_row):
        if h is not None and str(h).strip().upper() == name.upper():
            return i
    return None


def load_sheet_rates(wb, sheet_name, rate_column_name):
    """Read a sheet into {normalized_product_name: (rate, original_name)}."""
    if sheet_name not in wb.sheetnames:
        print(f"  ! WARNING: sheet '{sheet_name}' not found in workbook - skipping.")
        return {}
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    # Find the header row: the first row containing "PRODUCT NAME"
    header_idx = None
    for i, row in enumerate(rows):
        if row and any(c is not None and str(c).strip().upper() == "PRODUCT NAME" for c in row):
            header_idx = i
            break
    if header_idx is None:
        print(f"  ! WARNING: couldn't find header row in sheet '{sheet_name}' - skipping.")
        return {}
    header = rows[header_idx]
    name_col = find_column(header, "PRODUCT NAME")
    rate_col = find_column(header, rate_column_name)
    if rate_col is None:
        print(f"  ! WARNING: column '{rate_column_name}' not found in sheet '{sheet_name}' - skipping.")
        return {}

    out = {}
    for row in rows[header_idx + 1:]:
        if row is None or name_col >= len(row):
            continue
        name = row[name_col]
        if not name or not str(name).strip():
            continue
        rate = row[rate_col] if rate_col < len(row) else None
        if rate is None or not isinstance(rate, (int, float)):
            continue
        out[normalize(name)] = (float(rate), str(name).strip())
    return out


def load_rate_list(path):
    """Load the whole master rate list into one combined lookup:
    {family: {normalized_name: (rate, original_name)}}
    family is one of 'PVC_PIPE', 'PVC_FITTING', 'APVC', 'CPVC'."""
    print(f"Loading rate list: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    tables = {}
    for family, rule in SHEET_RULES.items():
        tables[family] = load_sheet_rates(wb, rule["sheet"], rule["column"])
        print(f"  {family}: {len(tables[family])} items loaded from '{rule['sheet']}' (column '{rule['column']}')")
    return tables


def resolve_rate_list_path(configured_path):
    if os.path.exists(configured_path):
        return configured_path
    print(f"Could not find '{configured_path}' in the current folder.")
    candidates = glob.glob("*.xlsx")
    candidates = [c for c in candidates if "RATE" in c.upper()]
    if candidates:
        print("Found a possible rate list file:")
        for c in candidates:
            print(f"  - {c}")
        use = input(f"Use '{candidates[0]}'? [Y/n] ").strip().lower()
        if use in ("", "y", "yes"):
            return candidates[0]
    typed = input("Type the full path to your rate list .xlsx file: ").strip()
    return typed


# =====================================================================
# Classification: which family / pipe-or-fitting does an item belong to?
# =====================================================================

def classify_family(item_name):
    """Return 'APVC', 'CPVC', or 'PVC' based on the item name. Checked in
    this order because 'APVC' and 'CPVC' both contain the letters 'PVC'."""
    n = normalize(item_name)
    if n.startswith("APVC") or " APVC" in n:
        return "APVC"
    if n.startswith("CPVC") or " CPVC" in n:
        return "CPVC"
    return "PVC"


def is_pipe(item_name):
    return "PIPE" in normalize(item_name)


def is_ten_kg(item_name):
    n = normalize(item_name)
    return any(k in n for k in TEN_KG_KEYWORDS)


def is_no_discount_item(item_name):
    n = normalize(item_name)
    return any(k in n for k in NO_DISCOUNT_KEYWORDS)


# =====================================================================
# Matching an order-line item name to a rate-list row
# =====================================================================

def apply_synonyms(norm_name):
    """Substitute known shorthand words/phrases with the rate-list's own
    wording, for matching purposes only."""
    out = norm_name
    for k, v in SYNONYMS.items():
        out = re.sub(rf"\b{re.escape(k)}\b", v, out)
    return out


def _token_dice(tokens_a, tokens_b):
    """Order-independent similarity: 2*|intersection| / (|A|+|B|), counting
    repeated tokens by multiset. This matters a lot here because bills
    write items as 'ELBOW 1I' (type then size) while the rate list writes
    '1I ELBOW AST' (size then type) - a plain character-by-character
    comparison scores that badly, token overlap doesn't."""
    if not tokens_a or not tokens_b:
        return 0.0
    a, b = list(tokens_a), list(tokens_b)
    inter = 0
    b_remaining = b[:]
    for t in a:
        if t in b_remaining:
            b_remaining.remove(t)
            inter += 1
    return 2 * inter / (len(a) + len(b))


def match_item(item_name, family_table):
    """Find the best match for item_name inside family_table
    ({normalized_name: (rate, original_name)}).
    Returns (rate, original_name, score, matched_via_synonym: bool) or
    (None, None, 0.0, False) if the table is empty.

    Scoring is token-based (word-overlap, order-independent) rather than
    plain character-sequence similarity, because product names in bills
    and in the rate list often list size/type in different order
    ('CPVC ELBOW OIII' vs 'CPVC OIII ELBOW AST')."""
    if not family_table:
        return None, None, 0.0, False

    norm = normalize(item_name)
    syn = apply_synonyms(norm)
    used_synonym = syn != norm

    # exact match first (post-synonym)
    if syn in family_table:
        rate, orig = family_table[syn]
        return rate, orig, 1.0, used_synonym

    my_tokens = syn.split()
    best_key, best_score = None, -1.0
    for key in family_table:
        score = _token_dice(my_tokens, key.split())
        if score > best_score:
            best_key, best_score = key, score

    if best_key is None:
        return None, None, 0.0, used_synonym
    rate, orig = family_table[best_key]
    return rate, orig, best_score, used_synonym


# =====================================================================
# Business rules: discount %, quantity conversion, special cases
# =====================================================================

def compute_quotation_row(item_name, order_qty, rate_tables):
    """Work out everything for one order line:
    returns dict with keys:
      name, family, order_qty, billing_qty, rate, discount_pct,
      rate_after_discount, amount, note
    """
    norm = normalize(item_name)
    note_parts = []

    # 1) Extra manually-supplied rate takes priority over the rate list
    if norm in EXTRA_RATES:
        rate, is_fitting_disc, extra_note = EXTRA_RATES[norm]
        note_parts.append(extra_note)
        family = classify_family(item_name)
        pipe = is_pipe(item_name)
        billing_qty = order_qty
        if pipe:
            billing_qty = order_qty / PIPE_QTY_DIVISOR.get(family, 1)
        if is_fitting_disc:
            discount = _fitting_discount_for_family(family)
        else:
            discount = 0.0
        rate_after = rate * (1 - discount)
        amount = billing_qty * rate_after
        return _row(item_name, family, order_qty, billing_qty, rate, discount,
                    rate_after, amount, "; ".join(note_parts))

    # 2) Otherwise, look it up in the rate list
    family = classify_family(item_name)
    pipe = is_pipe(item_name)
    ten_kg = pipe and is_ten_kg(item_name)

    if family == "PVC":
        # 10KG / NON-ISI pipes only exist in the "PVC" sheet (per-foot
        # rate), never in "PVC (2)" (which only lists the ISI/AST
        # full-pipe-price lines) - route them there instead.
        if pipe and not ten_kg:
            table = rate_tables["PVC_PIPE"]
        else:
            table = rate_tables["PVC_FITTING"]
    else:
        table = rate_tables[family]

    rate, matched_name, score, used_synonym = match_item(item_name, table)

    if rate is None:
        note_parts.append("NOT FOUND in rate list - please check item name / add manually")
        return _row(item_name, family, order_qty, order_qty, 0.0, 0.0, 0.0, 0.0,
                    "; ".join(note_parts))

    if matched_name and normalize(matched_name) != norm:
        note_parts.append(f"matched to: {matched_name}")
    if used_synonym:
        note_parts.append("matched via abbreviation")
    if score < CONFIDENCE_THRESHOLD:
        note_parts.append("LOW CONFIDENCE MATCH - please verify")

    if ten_kg:
        rate = rate * TEN_KG_RATE_MULTIPLIER
        note_parts.append(f"10KG pipe: rate x{TEN_KG_RATE_MULTIPLIER}")

    billing_qty = order_qty
    if pipe:
        billing_qty = order_qty / PIPE_QTY_DIVISOR.get(family, 1)

    # discount
    if is_no_discount_item(item_name) or ten_kg:
        discount = 0.0
        if not ten_kg:
            note_parts.append("no discount (excluded category)")
    elif pipe:
        if PIPE_DISCOUNT_REQUIRES_BRAND in norm:
            discount = DISCOUNT_CPVC_PIPE if family == "CPVC" else DISCOUNT_PVC_APVC_PIPE
        else:
            discount = 0.0
            note_parts.append("no discount (non-AST / unbranded pipe)")
    else:
        discount = _fitting_discount_for_family(family)

    rate_after = rate * (1 - discount)
    amount = billing_qty * rate_after

    return _row(item_name, family, order_qty, billing_qty, rate, discount,
                rate_after, amount, "; ".join(note_parts))


def _fitting_discount_for_family(family):
    return DISCOUNT_CPVC_FITTING if family == "CPVC" else DISCOUNT_PVC_APVC_FITTING


def _row(name, family, order_qty, billing_qty, rate, discount, rate_after, amount, note):
    return {
        "name": name, "family": family, "order_qty": order_qty,
        "billing_qty": billing_qty, "rate": rate, "discount_pct": discount,
        "rate_after_discount": rate_after, "amount": amount, "note": note,
    }


# =====================================================================
# Excel output: Order Quotation
# =====================================================================

def generate_quotation_excel(rows, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Order Quotation"
    fn = "Arial"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "ORDER QUOTATION"
    ws["A1"].font = Font(name=fn, bold=True, size=14)
    ws.merge_cells("A1:I1")

    headers = ["SR NO", "PRODUCT NAME", "ORDER QTY", "BILLING QTY", "RATE",
               "DISCOUNT %", "RATE AFTER DISCOUNT", "AMOUNT", "NOTE"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c, h)
        cell.font = Font(name=fn, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    plain_font = Font(name=fn)

    r = 4
    for i, row in enumerate(rows, 1):
        for col, val in [(1, i), (2, row["name"]), (3, row["order_qty"]),
                          (4, round(row["billing_qty"], 4)), (5, round(row["rate"], 4))]:
            cell = ws.cell(r, col, val)
            cell.border = border
            cell.font = plain_font

        dc = ws.cell(r, 6, row["discount_pct"])
        dc.number_format = "0.000%"
        dc.border = border
        dc.font = plain_font

        rc = ws.cell(r, 7, round(row["rate_after_discount"], 4))
        rc.border = border
        rc.font = plain_font

        # AMOUNT as a live formula: billing qty x rate after discount
        amt_cell = ws.cell(r, 8, f"=D{r}*G{r}")
        amt_cell.number_format = "0.00"
        amt_cell.border = border
        amt_cell.font = plain_font

        note_cell = ws.cell(r, 9, row["note"])
        note_cell.border = border
        note_cell.font = Font(name=fn, italic=True, size=9, color="808080")
        r += 1

    total_row = r
    tl_cell = ws.cell(total_row, 7, "GRAND TOTAL")
    tl_cell.font = Font(name=fn, bold=True)
    total_cell = ws.cell(total_row, 8, f"=SUM(H4:H{r - 1})")
    total_cell.number_format = "0.00"
    total_cell.font = Font(name=fn, bold=True)
    for c in range(1, 10):
        ws.cell(total_row, c).border = border

    # column widths
    widths = {"A": 6, "B": 34, "C": 11, "D": 12, "E": 10, "F": 11, "G": 16, "H": 12, "I": 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A4"

    wb.save(output_path)
    print(f"Saved: {output_path}")
    print("NOTE: open this file in Excel/LibreOffice once so it recalculates the")
    print("      AMOUNT and GRAND TOTAL formulas (openpyxl doesn't compute them).")


def read_item_list_csv(path):
    """CSV with two columns: item name, qty (no header, or a header row -
    both are handled)."""
    items = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            if not row or not row[0].strip():
                continue
            name = row[0].strip()
            if name.upper() in ("ITEM", "PRODUCT NAME", "NAME"):
                continue  # header row
            qty_raw = row[1].strip() if len(row) > 1 else ""
            try:
                qty = float(qty_raw)
            except ValueError:
                print(f"  ! Skipping row (couldn't read quantity): {row}")
                continue
            items.append((name, qty))
    return items


def read_item_list_interactive():
    print("\nType each item as:  Item Name, Qty")
    print("(e.g.  PVC PIPE 1I X 6KG ISI AST, 100 )")
    print("Press Enter on an empty line when you're done.\n")
    items = []
    while True:
        line = input("> ").strip()
        if not line:
            break
        if "," not in line:
            print("  ! Please use a comma between the name and the quantity. Try again.")
            continue
        name, qty_raw = line.rsplit(",", 1)
        try:
            qty = float(qty_raw.strip())
        except ValueError:
            print("  ! Couldn't read the quantity. Try again.")
            continue
        items.append((name.strip(), qty))
    return items


# =====================================================================
# PDF extraction: pull PVC/APVC/CPVC items + qty out of a sales bill
# =====================================================================

def extract_items_from_pdf(pdf_path, families=("PVC", "APVC", "CPVC")):
    """Very lightweight line-item extractor for the simple tabular bills
    this shop uses (Sr | Product Name | Qty | Rate | Amount). Returns a
    list of (item_name, qty) for lines whose product name belongs to one
    of `families`. The bill's own Rate/Amount columns are ignored - use
    compute_quotation_row() afterwards to fill correct rates.

    This is a best-effort text parser; always eyeball the result against
    the PDF, especially for multi-line product names."""
    try:
        import pdfplumber
    except ImportError:
        print("This feature needs the 'pdfplumber' package.")
        print("Install it with:  pip install pdfplumber")
        return []

    line_re = re.compile(
        r"^\s*(\d{1,3})\s+(.+?)\s+([\d,]+\.\d{3})\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})\s*$"
    )
    results = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.split("\n"):
                m = line_re.match(line)
                if not m:
                    continue
                sr, name, qty_str, rate_str, amount_str = m.groups()
                name = name.strip()
                fam = classify_family(name)
                # only keep it if the name actually starts with a family
                # keyword (avoids matching unrelated rows like brass items)
                n = normalize(name)
                if not any(n.startswith(f) for f in families):
                    continue
                try:
                    qty = float(qty_str.replace(",", ""))
                except ValueError:
                    continue
                results.append((name, qty))
    return results


# =====================================================================
# Update a price in the master rate list
# =====================================================================

def update_rate_list_price(path, sheet_name, item_contains, new_rate, column_name="RATE"):
    """Find every row in `sheet_name` whose PRODUCT NAME contains
    `item_contains` (case-insensitive) and overwrite the given column
    with `new_rate`. Saves the file in place. Returns the list of rows
    changed (original name, old rate)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")
        return []
    ws = wb[sheet_name]
    rows = list(ws.iter_rows())
    header_idx = None
    for i, row in enumerate(rows):
        vals = [c.value for c in row]
        if any(v is not None and str(v).strip().upper() == "PRODUCT NAME" for v in vals):
            header_idx = i
            break
    if header_idx is None:
        print("Couldn't find the header row (PRODUCT NAME) in this sheet.")
        return []

    header_vals = [c.value for c in rows[header_idx]]
    name_col = find_column(header_vals, "PRODUCT NAME")
    rate_col = find_column(header_vals, column_name)
    if rate_col is None:
        print(f"Column '{column_name}' not found. Header row is: {header_vals}")
        return []

    changed = []
    needle = item_contains.strip().upper()
    for row in rows[header_idx + 1:]:
        name_cell = row[name_col]
        if name_cell.value and needle in str(name_cell.value).strip().upper():
            rate_cell = row[rate_col]
            changed.append((name_cell.value, rate_cell.value))
            rate_cell.value = new_rate

    if changed:
        wb.save(path)
        print(f"Updated {len(changed)} row(s) in sheet '{sheet_name}', column '{column_name}':")
        for name, old in changed:
            print(f"  {name}: {old} -> {new_rate}")
    else:
        print(f"No rows found containing '{item_contains}' in sheet '{sheet_name}'.")
    return changed


# =====================================================================
# CLI menu
# =====================================================================

def menu_generate_quotation(rate_tables):
    print("\n--- Generate Order Quotation ---")
    print("1) Read items from a CSV file (Item Name, Qty per row)")
    print("2) Type items in one by one")
    choice = input("Choose 1 or 2: ").strip()

    if choice == "1":
        path = input("Path to your CSV file: ").strip()
        if not os.path.exists(path):
            print("File not found.")
            return
        items = read_item_list_csv(path)
    else:
        items = read_item_list_interactive()

    if not items:
        print("No items to process.")
        return

    rows = [compute_quotation_row(name, qty, rate_tables) for name, qty in items]

    grand_total = sum(r["amount"] for r in rows)
    print(f"\n{len(rows)} item(s) processed. Estimated grand total: {grand_total:,.2f}")
    not_found = [r for r in rows if r["rate"] == 0.0 and "NOT FOUND" in r["note"]]
    if not_found:
        print(f"! {len(not_found)} item(s) could not be matched - check the NOTE column in the output.")

    out_path = input("Output filename [Order_Quotation.xlsx]: ").strip() or "Order_Quotation.xlsx"
    generate_quotation_excel(rows, out_path)


def menu_extract_pdf(rate_tables):
    print("\n--- Extract items from a Sales Bill PDF ---")
    path = input("Path to the PDF: ").strip()
    if not os.path.exists(path):
        print("File not found.")
        return
    fam_in = input("Families to extract (comma-separated, default PVC,APVC,CPVC): ").strip()
    families = tuple(f.strip().upper() for f in fam_in.split(",")) if fam_in else ("PVC", "APVC", "CPVC")

    items = extract_items_from_pdf(path, families=families)
    if not items:
        print("No matching items found (or pdfplumber isn't installed).")
        return

    print(f"\nFound {len(items)} item(s):")
    for name, qty in items:
        print(f"  {name}  |  qty {qty}")

    fill = input("\nFill rates from the master rate list now? [Y/n] ").strip().lower()
    if fill in ("", "y", "yes"):
        rows = [compute_quotation_row(name, qty, rate_tables) for name, qty in items]
        out_path = input("Output filename [Extracted_Items.xlsx]: ").strip() or "Extracted_Items.xlsx"
        generate_quotation_excel(rows, out_path)
    else:
        out_path = input("Save the plain item/qty list to CSV as [Extracted_Items.csv]: ").strip() or "Extracted_Items.csv"
        with open(out_path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["Item Name", "Qty"])
            w.writerows(items)
        print(f"Saved: {out_path}")


def menu_update_rate():
    print("\n--- Update a rate in the master rate list ---")
    path = resolve_rate_list_path(RATE_LIST_PATH)
    if not path or not os.path.exists(path):
        print("File not found.")
        return
    wb = openpyxl.load_workbook(path, read_only=True)
    print(f"Sheets: {', '.join(wb.sheetnames)}")
    sheet = input("Sheet name to update: ").strip()
    item = input("Text to search for in PRODUCT NAME (matches any row containing it): ").strip()
    column = input("Column to update [RATE]: ").strip() or "RATE"
    new_rate_raw = input("New rate: ").strip()
    try:
        new_rate = float(new_rate_raw)
    except ValueError:
        print("That doesn't look like a number.")
        return
    confirm = input(f"About to set {column} = {new_rate} for every row in '{sheet}' containing "
                     f"'{item}'. Proceed? [y/N] ").strip().lower()
    if confirm != "y":
        print("Cancelled.")
        return
    update_rate_list_price(path, sheet, item, new_rate, column_name=column)


def main():
    print("=" * 60)
    print(" Salman's Pipe & Fitting Order Quotation App")
    print("=" * 60)

    rate_tables = None

    while True:
        print("\nWhat would you like to do?")
        print("  1) Generate an Order Quotation (priced + discounted)")
        print("  2) Extract items from a Sales Bill PDF and fill rates")
        print("  3) Update a rate in the master rate list")
        print("  4) Reload the rate list (after you've updated prices)")
        print("  0) Exit")
        choice = input("> ").strip()

        if choice in ("1", "2") and rate_tables is None:
            path = resolve_rate_list_path(RATE_LIST_PATH)
            if not path or not os.path.exists(path):
                print("Could not locate the rate list - try again.")
                continue
            rate_tables = load_rate_list(path)

        if choice == "1":
            menu_generate_quotation(rate_tables)
        elif choice == "2":
            menu_extract_pdf(rate_tables)
        elif choice == "3":
            menu_update_rate()
        elif choice == "4":
            path = resolve_rate_list_path(RATE_LIST_PATH)
            rate_tables = load_rate_list(path) if path else rate_tables
        elif choice == "0":
            print("Bye!")
            break
        else:
            print("Please choose one of the numbers above.")


if __name__ == "__main__":
    main()
